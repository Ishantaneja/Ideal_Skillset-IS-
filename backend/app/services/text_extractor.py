import io
import re
import logging
from typing import Union
import pypdf
import docx

logger = logging.getLogger("uvicorn.error")


class TextExtractor:
    """
    Extracts and normalizes text from PDF and DOCX resume documents.
    """

    @staticmethod
    def clean_text(raw_text: str) -> str:
        """
        Cleans and normalizes extracted resume text while preserving structural layout.
        """
        if not raw_text:
            return ""

        # Replace non-breaking spaces and non-standard whitespace
        text = raw_text.replace("\xa0", " ").replace("\r\n", "\n").replace("\r", "\n")

        # Strip unprintable control characters except \n and \t
        text = "".join(ch for ch in text if ch.isprintable() or ch in ("\n", "\t"))

        # Normalize multiple spaces into single space
        text = re.sub(r"[ \t]+", " ", text)

        # Condense 3+ newlines into 2 newlines
        text = re.sub(r"\n{3,}", "\n\n", text)

        return text.strip()

    @classmethod
    def extract_from_pdf(cls, source: Union[str, bytes]) -> str:
        """
        Extracts multi-page text from a PDF file path or byte stream using pypdf.
        """
        try:
            if isinstance(source, bytes):
                pdf_file = io.BytesIO(source)
            else:
                pdf_file = open(source, "rb")

            reader = pypdf.PdfReader(pdf_file)
            extracted_pages = []

            for idx, page in enumerate(reader.pages):
                try:
                    page_text = page.extract_text()
                    if page_text:
                        extracted_pages.append(page_text)
                except Exception as page_err:
                    logger.warning(f"Could not extract text from PDF page {idx + 1}: {page_err}")

            if not isinstance(source, bytes):
                pdf_file.close()

            full_text = "\n\n".join(extracted_pages)
            return cls.clean_text(full_text)
        except Exception as e:
            logger.error(f"PDF extraction error: {e}")
            raise ValueError(f"Failed to extract text from PDF: {str(e)}")

    @classmethod
    def extract_from_docx(cls, source: Union[str, bytes]) -> str:
        """
        Extracts paragraphs and table contents from a DOCX file using python-docx.
        """
        try:
            if isinstance(source, bytes):
                docx_file = io.BytesIO(source)
            else:
                docx_file = source

            doc = docx.Document(docx_file)
            text_blocks = []

            # 1. Paragraphs
            for p in doc.paragraphs:
                cleaned = p.text.strip()
                if cleaned:
                    text_blocks.append(cleaned)

            # 2. Table contents (e.g. skills tables, experience columns)
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        text_blocks.append(" | ".join(row_cells))

            full_text = "\n".join(text_blocks)
            return cls.clean_text(full_text)
        except Exception as e:
            logger.error(f"DOCX extraction error: {e}")
            raise ValueError(f"Failed to extract text from DOCX document: {str(e)}")

    @classmethod
    def extract_text(cls, file_path: str, file_type: str) -> str:
        """
        Unified extraction entrypoint for PDF and DOCX documents.
        """
        ft = file_type.lower().lstrip(".")
        if ft == "pdf":
            return cls.extract_from_pdf(file_path)
        elif ft in ("docx", "doc"):
            return cls.extract_from_docx(file_path)
        else:
            raise ValueError(f"Unsupported file type '{file_type}' for text extraction")


text_extractor = TextExtractor()

