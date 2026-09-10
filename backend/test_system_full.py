import io
import docx
from datetime import datetime, timezone, timedelta
from fastapi.testclient import TestClient
from app.main import app
from app.database.connection import mongo_manager
from app.services.otp_service import OTPService, _IN_MEMORY_SIGNUP_OTPS

client = TestClient(app)

def test_system_full_flow():
    print("===========================================================")
    print("   IDEAL SKILLSET — COMPREHENSIVE END-TO-END TEST SUITE   ")
    print("===========================================================")

    # 1. Health
    print("\n[1/12] Testing Health Endpoints...")
    res = client.get("/api/health")
    assert res.status_code == 200
    print("  [OK] Health API Status:", res.json()["status"])

    # 2. Direct Candidate Signup (No OTP)
    print("\n[2/12] Testing Direct Candidate Signup Flow (No OTP)...")
    candidate_email = f"test_candidate_{int(datetime.now().timestamp())}@idealskillset.com"
    res = client.post("/api/auth/signup", json={
        "name": "Elena Rostova",
        "email": candidate_email,
        "password": "CandidatePassword123!",
        "target_role": "Data Analyst"
    })
    assert res.status_code == 201, f"Direct Signup failed: {res.text}"
    user_token = res.json()["access_token"]
    headers = {"Authorization": f"Bearer {user_token}"}
    print("  [OK] Direct candidate registration successful. Account created in User_data and JWT token issued.")

    # 3. Authentication & Profile
    print("\n[3/12] Testing Candidate Login & Profile...")
    res = client.post("/api/auth/login", json={"email": candidate_email, "password": "CandidatePassword123!"})
    assert res.status_code == 200
    assert "access_token" in res.json()

    res = client.get("/api/users/me", headers=headers)
    assert res.status_code == 200
    print("  [OK] Candidate Profile retrieved for:", res.json()["name"])

    # Update profile
    res = client.put("/api/users/me", json={
        "target_role": "Senior Data Analyst",
        "github_url": "https://github.com/elenarostova/analytics",
        "skills": ["SQL", "Python", "Power BI", "Excel"]
    }, headers=headers)
    assert res.status_code == 200
    print("  [OK] Candidate Profile updated with GitHub URL & target role.")

    # 4. Admin Auth
    print("\n[4/12] Testing Admin Authentication...")
    res = client.post("/api/auth/admin/login", json={"email": "admin@idealskillset.com", "password": "Admin1234!"})
    assert res.status_code == 200
    admin_token = res.json()["access_token"]
    admin_headers = {"Authorization": f"Bearer {admin_token}"}
    print("  [OK] Admin logged in successfully with admin privileges.")

    # 5. Resume Upload & Parsing
    print("\n[5/12] Testing Resume Upload & Extraction...")
    doc = docx.Document()
    doc.add_heading("Elena Rostova - Senior Data Analyst", level=1)
    doc.add_paragraph("Summary: Experienced analyst proficient in SQL, Python, Excel, and Power BI. GitHub: https://github.com/elenarostova/analytics")
    doc.add_paragraph("Experience: Data Analyst at Apex Global (2022-2025). Designed automated SQL pipelines processing 500k+ rows. Built interactive Power BI dashboards.")
    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    res = client.post("/api/resumes/upload", files={"file": ("elena_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}, headers=headers)
    assert res.status_code == 201
    resume_id = res.json()["id"]
    skills_count = len(res.json().get("parsed_data", {}).get("skills", []))
    print(f"  [OK] Resume uploaded ({resume_id}): Extracted {skills_count} precision skills.")

    # 6. Job Description Analysis
    print("\n[6/12] Testing Job Description Analyzer...")
    job_text = """
Job Title: Senior Data Analyst
Company: Enterprise Analytics Corp
Requirements:
• 3+ years of experience with Advanced SQL, query optimization, and window functions.
• Strong hands-on expertise in Power BI and DAX measures.
• Proficiency with Python (Pandas, NumPy) for data processing.
Preferred:
• Experience with Tableau and Docker.
"""
    res = client.post("/api/jobs/analyze", json={"text": job_text}, headers=headers)
    assert res.status_code == 201
    job_id = res.json()["id"]
    print(f"  [OK] Job Description analyzed ({job_id}): Extracted requirements & required skills.")

    # 7. Explainable ATS Scoring & Simulation
    print("\n[7/12] Testing Explainable ATS Matcher & What-If Simulation...")
    res = client.post("/api/ats/analyze", json={"resume_id": resume_id, "job_id": job_id}, headers=headers)
    assert res.status_code == 201
    ats_id = res.json()["id"]
    ats_score = res.json()["score"]
    print(f"  [OK] ATS Match evaluated ({ats_id}): Score = {ats_score}%")

    # What-if simulation
    res = client.post("/api/ats/simulate", json={"resume_id": resume_id, "job_id": job_id, "added_skills": ["Tableau", "Docker"]}, headers=headers)
    assert res.status_code == 200
    sim_score = res.json()["simulated_score"]
    score_gain = res.json()["score_gain"]
    print(f"  [OK] What-If Simulation: Score increased to {sim_score}% (+{score_gain}%)")

    # 8. Skill Gap Analysis & Priority Ranking
    print("\n[8/12] Testing Intelligent Skill Gap Analysis...")
    res = client.post("/api/skill-gaps/analyze", json={"resume_id": resume_id, "job_id": job_id}, headers=headers)
    assert res.status_code == 201
    sg_id = res.json()["id"]
    top_gap = res.json()["skills"][0]["skill"] if res.json()["skills"] else "None"
    print(f"  [OK] Skill Gap analyzed ({sg_id}): Top priority gap identified as '{top_gap}'.")

    # 9. Personalized Career Roadmap
    print("\n[9/12] Testing Personalized Multi-Week Career Roadmap...")
    res = client.post("/api/roadmaps/generate", json={"resume_id": resume_id, "job_id": job_id, "duration_weeks": 4}, headers=headers)
    assert res.status_code == 201
    rm_id = res.json()["id"]
    first_task_id = res.json()["weeks"][0]["tasks"][0]["id"]

    # Patch task
    res = client.patch(f"/api/roadmaps/{rm_id}/tasks/{first_task_id}", json={"status": "completed"}, headers=headers)
    assert res.status_code == 200
    progress = res.json()["overall_progress"]
    print(f"  [OK] Roadmap generated ({rm_id}) & task completed: Progress updated to {progress}%.")

    # 10. 5D AI Readiness Twin
    print("\n[10/12] Testing 5-Dimensional AI Readiness Twin...")
    res = client.post("/api/readiness/analyze", json={"resume_id": resume_id, "job_id": job_id, "github_url": "https://github.com/elenarostova/analytics"}, headers=headers)
    assert res.status_code == 201
    rt_id = res.json()["id"]
    readiness_score = res.json()["overall_readiness_score"]
    verdict = res.json()["verdict"]["verdict_label"]
    print(f"  [OK] Readiness Twin evaluated ({rt_id}): Score = {readiness_score}%, Verdict = '{verdict}'.")

    # 11. Role-Adaptive Assessments & Interviews
    print("\n[11/12] Testing Role-Adaptive Assessments & Mock Interviews...")
    res = client.get("/api/assessment/challenges", headers=headers)
    assert res.status_code == 200
    print("  [OK] Adaptive Assessments returned for:", res.json()["target_role"])

    res = client.get("/api/interview/questions", headers=headers)
    assert res.status_code == 200
    print("  [OK] Adaptive Mock Interview Questions returned for:", res.json()["target_role"])

    # 12. Admin Endpoints & Security RBAC
    print("\n[12/12] Testing Admin Analytics & RBAC Security...")
    res = client.get("/api/admin/overview", headers=admin_headers)
    assert res.status_code == 200
    print("  [OK] Admin Overview metrics accessible to admin.")

    res = client.get("/api/admin/users", headers=admin_headers)
    assert res.status_code == 200
    print("  [OK] Admin Users directory count:", res.json()["total"])

    # Security test: Candidate gets 403 Forbidden on Admin endpoint
    res = client.get("/api/admin/users", headers=headers)
    assert res.status_code == 403
    print("  [OK] Security: Candidate rejected with 403 Forbidden on admin routes.")

    print("\n===========================================================")
    print("       ALL 12/12 INTEGRATION TEST MODULES PASSED! [100%]    ")
    print("===========================================================")

if __name__ == "__main__":
    test_system_full_flow()
